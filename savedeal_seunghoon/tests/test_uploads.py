import io

from openpyxl import Workbook, load_workbook

from config import Config
from repositories.customer_repository import CustomerRepository
from repositories.inventory_repository import InventoryRepository

RES_HEADERS = ["예약번호", "고객번호*", "고객명", "연락처", "매장코드*", "모델*", "색상*", "용량*",
               "가입유형*", "기존 통신사", "희망 수령일*", "메모"]
GOOD_ROW = ["R3001", "C001", "김지수", "010-****-1024", "S001", "GalaxyZ Fold6", "Black", "256GB",
            "번호이동", "LG U+", "2026-10-10", "블랙도 괜찮다고 함"]


def xlsx(headers, rows) -> bytes:
    book = Workbook()
    sheet = book.active
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def upload(client, kind, content, filename="list.xlsx"):
    return client.post(
        "/api/uploads",
        data={"kind": kind, "file": (io.BytesIO(content), filename)},
        content_type="multipart/form-data",
    )


def test_preview_separates_valid_and_error_rows(client):
    rows = [
        GOOD_ROW,
        ["", "C002", "", "010-1234-5678", "S001", "GalaxyZ Fold6", "Black", "256GB", "신규가입", "", "2026-10-10", ""],
        ["", "C003", "", "", "S999", "Unknown", "Red", "1TB", "번호이동", "", "10/10", ""],
    ]
    data = upload(client, "reservations", xlsx(RES_HEADERS, rows)).get_json()["data"]

    assert (data["total_rows"], data["valid_rows"], data["error_rows"]) == (3, 1, 2)
    messages = " ".join(m for e in data["errors"] for m in e["messages"])
    assert "가운데를 가린 형식" in messages  # 실제 번호 차단
    assert "매장코드" in messages and "단말 목록" in messages
    assert "기존 통신사" in messages and "희망 수령일" in messages
    assert [e["row"] for e in data["errors"]] == [3, 4]
    assert data["preview"][0]["line_type"] == "MNP"
    assert data["preview"][0]["previous_carrier"] == "LGU"


def test_commit_creates_reservations_with_precheck(client):
    batch = upload(client, "reservations", xlsx(RES_HEADERS, [GOOD_ROW])).get_json()["data"]

    done = client.post(f"/api/uploads/{batch['batch_id']}/commit").get_json()["data"]
    assert done["status"] == "COMMITTED"
    assert done["result"]["reservation_ids"] == ["R3001"]
    assert "high_risk" in done["result"]

    detail = client.get("/api/reservations/R3001").get_json()["data"]
    assert detail["memo"] == "블랙도 괜찮다고 함"
    assert detail["carrier_change"]["label"] == "LG U+ → KT"
    assert any("엑셀 업로드" in h["description"] for h in detail["history"])


def test_commit_twice_is_rejected(client):
    batch = upload(client, "reservations", xlsx(RES_HEADERS, [GOOD_ROW])).get_json()["data"]
    client.post(f"/api/uploads/{batch['batch_id']}/commit")
    assert client.post(f"/api/uploads/{batch['batch_id']}/commit").status_code == 409


def test_existing_reservation_id_is_an_error(client):
    row = ["R2001"] + GOOD_ROW[1:]
    data = upload(client, "reservations", xlsx(RES_HEADERS, [row])).get_json()["data"]
    assert data["valid_rows"] == 0
    assert "이미 등록된 예약번호" in data["errors"][0]["messages"][0]


def test_customer_upload_upserts_and_affects_precheck(client, db_path):
    headers = ["고객번호*", "고객명*", "연락처*", "본인인증*", "필요서류", "제출서류", "미납*", "할부한도*"]
    rows = [["C001", "김지수", "010-****-1024", "예", "신분증", "신분증", "예", "2,000,000"],
            ["C301", "새고객", "010-****-9999", "아니오", "", "", "아니오", "1500000"]]
    batch = upload(client, "customers", xlsx(headers, rows)).get_json()["data"]
    done = client.post(f"/api/uploads/{batch['batch_id']}/commit").get_json()["data"]

    assert (done["result"]["created"], done["result"]["updated"]) == (1, 1)
    repo = CustomerRepository(Config.DATA_DIR, db_path)
    assert repo.find_by_id("C001")["overdue_payment"] is True
    assert repo.find_by_id("C301")["identity_verified"] is False


def test_inventory_upload_csv_with_korean_excel_encoding(client, db_path):
    csv_text = "SKU,매장코드,모델,색상,용량,재고수량,입고 예정일\nSKU001,S001,GalaxyZ Fold6,Black,256GB,0,2026-10-20\n"
    batch = upload(client, "inventory", csv_text.encode("cp949"), "stock.csv").get_json()["data"]
    client.post(f"/api/uploads/{batch['batch_id']}/commit")

    item = InventoryRepository(Config.DATA_DIR, db_path).find_by_id("SKU001")
    assert item["quantity_on_hand"] == 0
    assert item["expected_restock_date"] == "2026-10-20"


def test_cancel_and_history(client):
    batch = upload(client, "reservations", xlsx(RES_HEADERS, [GOOD_ROW])).get_json()["data"]
    client.post(f"/api/uploads/{batch['batch_id']}/cancel")

    items = client.get("/api/uploads").get_json()["data"]["items"]
    assert items[0]["status"] == "CANCELLED"
    assert client.get("/api/reservations/R3001").status_code == 404


def test_file_level_errors(client):
    assert upload(client, "reservations", b"hello", "list.txt").status_code == 400
    missing = upload(client, "reservations", xlsx(["고객번호"], [["C001"]])).get_json()
    assert "필수 열이 없습니다" in missing["error"]["message"]
    assert upload(client, "unknown", xlsx(RES_HEADERS, [GOOD_ROW])).status_code == 400


def test_template_download_has_headers(client):
    response = client.get("/api/uploads/template/reservations")
    assert response.status_code == 200
    sheet = load_workbook(io.BytesIO(response.data)).worksheets[0]
    headers = [cell.value for cell in sheet[1]]
    assert "고객번호*" in headers and "메모" in headers


def test_upload_page_renders(client):
    body = client.get("/savedeal/upload").get_data(as_text=True)
    assert 'id="upload-file"' in body and "사전예약 명단" in body
