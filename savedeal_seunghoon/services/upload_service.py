"""엑셀(.xlsx)·CSV 일괄 업로드: 읽기 → 행별 검증 → 미리보기 → 확정(저장).

- 원본 파일은 저장하지 않는다. 검증을 통과한 행만 확정 전까지 upload_batches 에 잠시 두고, 확정하면 지운다.
- 연락처는 가운데를 가린 형식(010-****-1234)만 받는다. 교육용 가상 데이터만 다루기 위한 장치다.
"""
import csv
import io
import json
import re
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

from db.connection import connect, ensure_database
from repositories.customer_repository import CustomerRepository
from repositories.device_repository import DeviceRepository
from repositories.inventory_repository import InventoryRepository
from repositories.reservation_repository import ReservationRepository
from repositories.store_repository import StoreRepository

MAX_FILE_BYTES = 5 * 1024 * 1024
MAX_ROWS = 1000
PREVIEW_ROWS = 20

LINE_TYPE_ALIASES = {"NEW": "NEW", "신규": "NEW", "신규가입": "NEW", "MNP": "MNP", "번호이동": "MNP",
                     "CHANGE": "CHANGE", "기변": "CHANGE", "기기변경": "CHANGE"}
CARRIER_ALIASES = {"SKT": "SKT", "SK텔레콤": "SKT", "SK": "SKT", "LGU": "LGU", "LG U+": "LGU", "LGU+": "LGU",
                   "LG유플러스": "LGU", "MVNO": "MVNO", "알뜰폰": "MVNO"}
YES = {"예", "Y", "YES", "O", "TRUE", "1", "완료"}
NO = {"아니오", "아니요", "N", "NO", "X", "FALSE", "0", "미완료"}
MASKED_PHONE = re.compile(r"^01[0-9]-\*{3,4}-\d{4}$")

# kind → 열 정의: (키, 엑셀 머리글, 필수 여부, 예시 값)
KINDS = {
    "reservations": {
        "label": "사전예약 명단",
        "description": "신규 예약을 한 번에 등록합니다. 등록하면 사전검증·문제 감지·해결책 생성이 자동으로 진행됩니다.",
        "columns": [
            ("reservation_id", "예약번호", False, "R3001"),
            ("customer_id", "고객번호", True, "C001"),
            ("customer_name", "고객명", False, "김지수"),
            ("customer_phone", "연락처", False, "010-****-1024"),
            ("store_id", "매장코드", True, "S001"),
            ("model", "모델", True, "Galaxy Z Fold8"),
            ("color", "색상", True, "Black"),
            ("storage", "용량", True, "256GB"),
            ("line_type", "가입유형", True, "번호이동"),
            ("previous_carrier", "기존 통신사", False, "LG U+"),
            ("desired_activation_date", "희망 수령일", True, "2026-10-10"),
            ("memo", "메모", False, "블랙도 괜찮다고 함, 서류는 월요일에 가져온다고 함"),
        ],
    },
    "customers": {
        "label": "고객 정보",
        "description": "고객 기준정보를 등록하거나 갱신합니다. 같은 고객번호가 있으면 갱신합니다.",
        "columns": [
            ("customer_id", "고객번호", True, "C201"),
            ("name", "고객명", True, "홍길동"),
            ("phone", "연락처", True, "010-****-5678"),
            ("identity_verified", "본인인증", True, "예"),
            ("required_documents", "필요서류", False, "신분증, 가족관계증명서"),
            ("submitted_documents", "제출서류", False, "신분증"),
            ("overdue_payment", "미납", True, "아니오"),
            ("installment_limit", "할부한도", True, "2000000"),
            ("existing_lines_count", "보유 회선", False, "1"),
            ("max_lines_allowed", "최대 회선", False, "5"),
        ],
    },
    "inventory": {
        "label": "재고 현황",
        "description": "매장별 재고를 등록하거나 갱신합니다. 같은 SKU가 있으면 갱신합니다.",
        "columns": [
            ("sku", "SKU", True, "SKU101"),
            ("store_id", "매장코드", True, "S001"),
            ("model", "모델", True, "Galaxy Z Fold8"),
            ("color", "색상", True, "Silver"),
            ("storage", "용량", True, "512GB"),
            ("quantity_on_hand", "재고수량", True, "3"),
            ("expected_restock_date", "입고 예정일", False, "2026-10-15"),
        ],
    },
}


class UploadError(Exception):
    """파일 전체를 처리할 수 없을 때 (형식, 크기, 머리글 누락 등)."""


def _clean_header(text) -> str:
    return re.sub(r"[\s*]|\(선택\)|\(필수\)", "", str(text or ""))


def _cell_text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def read_table(filename: str, content: bytes) -> list[dict]:
    """첫 시트(또는 CSV)를 {머리글: 값} 행 목록으로 읽는다. 행 번호(_row)는 엑셀 기준."""
    if len(content) > MAX_FILE_BYTES:
        raise UploadError("파일이 너무 큽니다 (최대 5MB).")
    suffix = Path(filename or "").suffix.lower()
    if suffix == ".xlsx":
        try:
            sheet = load_workbook(io.BytesIO(content), read_only=True, data_only=True).worksheets[0]
        except Exception as exc:  # 손상되었거나 엑셀이 아닌 파일
            raise UploadError("엑셀 파일을 읽지 못했습니다. .xlsx 형식인지 확인해 주세요.") from exc
        rows = [[_cell_text(v) for v in row] for row in sheet.iter_rows(values_only=True)]
    elif suffix == ".csv":
        for encoding in ("utf-8-sig", "cp949"):
            try:
                text = content.decode(encoding)
                break
            except UnicodeDecodeError:
                continue
        else:
            raise UploadError("CSV 글자 인코딩을 알 수 없습니다. UTF-8 또는 엑셀 CSV로 저장해 주세요.")
        rows = [[cell.strip() for cell in row] for row in csv.reader(io.StringIO(text))]
    else:
        raise UploadError("엑셀(.xlsx) 또는 CSV(.csv) 파일만 올릴 수 있습니다.")

    header_index = next((i for i, row in enumerate(rows) if any(cell for cell in row)), None)
    if header_index is None:
        raise UploadError("빈 파일입니다.")
    headers = [_clean_header(cell) for cell in rows[header_index]]
    records = []
    for offset, row in enumerate(rows[header_index + 1:], start=header_index + 2):
        if not any(row):
            continue
        record = {headers[i]: row[i] if i < len(row) else "" for i in range(len(headers)) if headers[i]}
        record["_row"] = offset
        records.append(record)
    if len(records) > MAX_ROWS:
        raise UploadError(f"한 번에 {MAX_ROWS}행까지 올릴 수 있습니다.")
    return records


def _parse_date(text: str) -> str | None:
    for fmt in ("%Y-%m-%d", "%Y/%m/%d", "%Y.%m.%d", "%Y%m%d"):
        try:
            return datetime.strptime(text, fmt).date().isoformat()
        except ValueError:
            continue
    return None


def _parse_bool(text: str) -> bool | None:
    upper = text.strip().upper()
    if upper in YES:
        return True
    if upper in NO:
        return False
    return None


def _parse_int(text: str) -> int | None:
    cleaned = text.replace(",", "").replace("원", "").strip()
    return int(cleaned) if re.fullmatch(r"\d+", cleaned) else None


def _split_list(text: str) -> list[str]:
    return [part.strip() for part in re.split(r"[,/;]", text) if part.strip()]


class UploadService:
    def __init__(self, db_path, data_dir: Path, clock=datetime.now):
        ensure_database(db_path)
        self.db_path = db_path
        self.data_dir = data_dir
        self.clock = clock
        self.customer_repo = CustomerRepository(data_dir, db_path)
        self.inventory_repo = InventoryRepository(data_dir, db_path)
        self.reservation_repo = ReservationRepository(db_path)
        self.store_ids = {store["store_id"] for store in StoreRepository(data_dir).load_all()}
        self.catalog = {(d["model"], d["color"], d["storage"]) for d in DeviceRepository(data_dir).load_all()}

    def _now(self) -> str:
        return self.clock().isoformat(timespec="seconds")

    # ── 양식 ───────────────────────────────────────────────────────

    def template(self, kind: str) -> bytes:
        spec = KINDS[kind]
        book = Workbook()
        sheet = book.active
        sheet.title = spec["label"]
        header_fill = PatternFill("solid", fgColor="F4F6F9")
        for index, (_, header, required, example) in enumerate(spec["columns"], start=1):
            cell = sheet.cell(row=1, column=index, value=f"{header}*" if required else header)
            cell.font = Font(bold=True, color="D42A37" if required else "151A23")
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center")
            sheet.cell(row=2, column=index, value=example)
            sheet.column_dimensions[cell.column_letter].width = max(12, len(header) * 2 + 6, len(example) + 4)
        guide = book.create_sheet("작성 안내")
        lines = [
            f"[{spec['label']}] {spec['description']}",
            "* 표시는 필수 열입니다. 2행은 예시이니 지우고 작성하세요.",
            "개인정보 보호를 위해 연락처는 가운데 4자리를 가린 형식(010-****-1234)만 받습니다.",
            "날짜는 2026-10-10 형식으로 입력하세요.",
            "매장코드: " + ", ".join(sorted(self.store_ids)),
        ]
        if kind != "customers":
            lines.append("등록 가능한 단말: " + " / ".join(f"{m} {c} {s}" for m, c, s in sorted(self.catalog)))
        for row, line in enumerate(lines, start=1):
            guide.cell(row=row, column=1, value=line)
        guide.column_dimensions["A"].width = 120
        buffer = io.BytesIO()
        book.save(buffer)
        return buffer.getvalue()

    # ── 검증 ───────────────────────────────────────────────────────

    def _validate_row(self, kind: str, raw: dict, seen: set) -> tuple[dict, list[str]]:
        spec = KINDS[kind]
        value = {key: raw.get(_clean_header(header), "").strip() for key, header, _, _ in spec["columns"]}
        errors = [f"'{header}' 값이 비어 있습니다." for key, header, required, _ in spec["columns"]
                  if required and not value[key]]

        if value.get("store_id") and value["store_id"] not in self.store_ids:
            errors.append(f"매장코드 '{value['store_id']}'를 찾을 수 없습니다.")
        if kind in ("reservations", "inventory") and value.get("model") and value.get("color") and value.get("storage"):
            if (value["model"], value["color"], value["storage"]) not in self.catalog:
                errors.append(f"단말 목록에 없는 모델·색상·용량입니다: {value['model']} {value['color']} {value['storage']}")

        phone_key = "customer_phone" if kind == "reservations" else "phone"
        if value.get(phone_key) and not MASKED_PHONE.match(value[phone_key]):
            errors.append("개인정보 보호를 위해 연락처는 가운데를 가린 형식(010-****-1234)만 받습니다.")

        if kind == "reservations":
            if value["customer_id"] and not re.fullmatch(r"C\d+", value["customer_id"]):
                errors.append("고객번호는 C001 같은 형식이어야 합니다.")
            if value["reservation_id"]:
                if not re.fullmatch(r"R\d+", value["reservation_id"]):
                    errors.append("예약번호는 R3001 같은 형식이어야 합니다.")
                elif value["reservation_id"] in seen:
                    errors.append("파일 안에 같은 예약번호가 두 번 있습니다.")
                elif self.reservation_repo.find_by_id(value["reservation_id"]):
                    errors.append(f"이미 등록된 예약번호입니다: {value['reservation_id']}")
                seen.add(value["reservation_id"])
            line_type = LINE_TYPE_ALIASES.get(value["line_type"].upper(), LINE_TYPE_ALIASES.get(value["line_type"]))
            if value["line_type"] and not line_type:
                errors.append("가입유형은 신규가입 / 번호이동 / 기기변경 중 하나여야 합니다.")
            value["line_type"] = line_type
            carrier = CARRIER_ALIASES.get(value["previous_carrier"].upper(), CARRIER_ALIASES.get(value["previous_carrier"]))
            if value["previous_carrier"] and not carrier:
                errors.append("기존 통신사는 SK텔레콤 / LG U+ / 알뜰폰 중 하나여야 합니다.")
            if line_type == "MNP" and not carrier and not value["previous_carrier"]:
                errors.append("번호이동은 기존 통신사를 입력해야 합니다.")
            value["previous_carrier"] = carrier if line_type == "MNP" else None
            if value["desired_activation_date"]:
                parsed = _parse_date(value["desired_activation_date"])
                if not parsed:
                    errors.append("희망 수령일은 2026-10-10 형식이어야 합니다.")
                value["desired_activation_date"] = parsed
            if len(value["memo"]) > 300:
                errors.append("메모는 300자까지 입력할 수 있습니다.")

        elif kind == "customers":
            if value["customer_id"] and not re.fullmatch(r"C\d+", value["customer_id"]):
                errors.append("고객번호는 C001 같은 형식이어야 합니다.")
            if value["customer_id"] in seen:
                errors.append("파일 안에 같은 고객번호가 두 번 있습니다.")
            seen.add(value["customer_id"])
            for key, header in (("identity_verified", "본인인증"), ("overdue_payment", "미납")):
                parsed = _parse_bool(value[key]) if value[key] else None
                if value[key] and parsed is None:
                    errors.append(f"'{header}'는 예 / 아니오로 입력해야 합니다.")
                value[key] = parsed
            for key, header, default in (("installment_limit", "할부한도", None),
                                         ("existing_lines_count", "보유 회선", 1),
                                         ("max_lines_allowed", "최대 회선", 5)):
                if value[key]:
                    parsed = _parse_int(value[key])
                    if parsed is None:
                        errors.append(f"'{header}'는 숫자여야 합니다.")
                    value[key] = parsed
                else:
                    value[key] = default
            value["required_documents"] = _split_list(value["required_documents"]) or ["신분증"]
            value["submitted_documents"] = _split_list(value["submitted_documents"])

        elif kind == "inventory":
            if value["sku"] in seen:
                errors.append("파일 안에 같은 SKU가 두 번 있습니다.")
            seen.add(value["sku"])
            quantity = _parse_int(value["quantity_on_hand"]) if value["quantity_on_hand"] else None
            if value["quantity_on_hand"] and quantity is None:
                errors.append("'재고수량'은 0 이상의 숫자여야 합니다.")
            value["quantity_on_hand"] = quantity
            if value["expected_restock_date"]:
                parsed = _parse_date(value["expected_restock_date"])
                if not parsed:
                    errors.append("입고 예정일은 2026-10-10 형식이어야 합니다.")
                value["expected_restock_date"] = parsed
            else:
                value["expected_restock_date"] = None

        return value, errors

    def preview(self, kind: str, filename: str, content: bytes) -> dict:
        if kind not in KINDS:
            raise UploadError("업로드 종류가 올바르지 않습니다.")
        records = read_table(filename, content)
        if not records:
            raise UploadError("데이터 행이 없습니다. 2행부터 내용을 입력해 주세요.")
        headers = set(records[0]) - {"_row"}
        missing = [header for _, header, required, _ in KINDS[kind]["columns"]
                   if required and _clean_header(header) not in headers]
        if missing:
            raise UploadError(f"필수 열이 없습니다: {', '.join(missing)}. 양식을 내려받아 열 이름을 맞춰 주세요.")

        valid, errors, seen = [], [], set()
        for raw in records:
            value, row_errors = self._validate_row(kind, raw, seen)
            if row_errors:
                errors.append({"row": raw["_row"], "messages": row_errors})
            else:
                valid.append({**value, "_row": raw["_row"]})

        now = self._now()
        with connect(self.db_path) as conn:
            cursor = conn.execute(
                """INSERT INTO upload_batches (kind, filename, status, total_rows, valid_rows, error_rows, rows, errors, created_at)
                   VALUES (?, ?, 'PREVIEW', ?, ?, ?, ?, ?, ?)""",
                (kind, Path(filename).name, len(records), len(valid), len(errors),
                 json.dumps(valid, ensure_ascii=False), json.dumps(errors, ensure_ascii=False), now),
            )
            batch_id = cursor.lastrowid
        return self.get_batch(batch_id, include_preview=True)

    # ── 조회 ───────────────────────────────────────────────────────

    def get_batch(self, batch_id: int, include_preview: bool = False) -> dict | None:
        with connect(self.db_path) as conn:
            row = conn.execute("SELECT * FROM upload_batches WHERE batch_id = ?", (batch_id,)).fetchone()
        if row is None:
            return None
        batch = dict(row)
        spec = KINDS[batch["kind"]]
        result = {
            "batch_id": batch["batch_id"],
            "kind": batch["kind"],
            "kind_label": spec["label"],
            "filename": batch["filename"],
            "status": batch["status"],
            "total_rows": batch["total_rows"],
            "valid_rows": batch["valid_rows"],
            "error_rows": batch["error_rows"],
            "errors": json.loads(batch["errors"]),
            "result": json.loads(batch["result"]) if batch["result"] else None,
            "created_at": batch["created_at"],
            "committed_at": batch["committed_at"],
        }
        if include_preview:
            rows = json.loads(batch["rows"])[:PREVIEW_ROWS]
            result["columns"] = [{"key": key, "label": header} for key, header, _, _ in spec["columns"]]
            result["preview"] = rows
        return result

    def list_batches(self, limit: int = 10) -> list[dict]:
        with connect(self.db_path) as conn:
            ids = [r[0] for r in conn.execute("SELECT batch_id FROM upload_batches ORDER BY batch_id DESC LIMIT ?", (limit,))]
        return [self.get_batch(batch_id) for batch_id in ids]

    # ── 확정 ───────────────────────────────────────────────────────

    def cancel(self, batch_id: int) -> dict:
        batch = self._pending(batch_id)
        with connect(self.db_path) as conn:
            conn.execute("UPDATE upload_batches SET status = 'CANCELLED', rows = '[]' WHERE batch_id = ?", (batch["batch_id"],))
        return self.get_batch(batch_id)

    def _pending(self, batch_id: int) -> dict:
        with connect(self.db_path) as conn:
            row = conn.execute("SELECT * FROM upload_batches WHERE batch_id = ?", (batch_id,)).fetchone()
        if row is None:
            raise LookupError("업로드를 찾을 수 없습니다.")
        if row["status"] != "PREVIEW":
            raise ValueError("이미 처리된 업로드입니다.")
        return dict(row)

    def commit(self, batch_id: int, reservation_service, dashboard_service=None) -> dict:
        """검증된 행을 저장한다. 예약은 reservation_service.create 로 등록해 사전검증·해결책 생성까지 진행한다."""
        batch = self._pending(batch_id)
        rows = json.loads(batch["rows"])
        now = self._now()
        result = {"created": 0, "updated": 0, "reservation_ids": []}

        if batch["kind"] == "customers":
            for row in rows:
                exists = self.customer_repo.find_by_id(row["customer_id"]) is not None
                fields = {key: row[key] for key, *_ in KINDS["customers"]["columns"]}
                self.customer_repo.upsert({**fields, "verification_method": "upload", "updated_at": now})
                result["updated" if exists else "created"] += 1
        elif batch["kind"] == "inventory":
            for row in rows:
                exists = self.inventory_repo.find_by_id(row["sku"]) is not None
                fields = {key: row[key] for key, *_ in KINDS["inventory"]["columns"]}
                self.inventory_repo.upsert({**fields, "updated_at": now})
                result["updated" if exists else "created"] += 1
        else:
            for row in rows:
                payload = {
                    "reservation_id": row["reservation_id"] or None,
                    "customer_id": row["customer_id"],
                    "customer_name": row["customer_name"] or None,
                    "customer_phone": row["customer_phone"] or None,
                    "store_id": row["store_id"],
                    "device": {"model": row["model"], "color": row["color"], "storage": row["storage"]},
                    "line_type": row["line_type"],
                    "previous_carrier": row["previous_carrier"],
                    "desired_activation_date": row["desired_activation_date"],
                    "memo": row["memo"] or None,
                }
                reservation_id, _ = reservation_service.create(payload, source="upload")
                result["reservation_ids"].append(reservation_id)
                result["created"] += 1

        if dashboard_service is not None and result["reservation_ids"]:
            details = [dashboard_service.get_detail(rid) for rid in result["reservation_ids"]]
            result["high_risk"] = sum(1 for d in details if d["risk_level"] == "high")
            result["due_today"] = sum(1 for d in details if d["is_due_today"])
            result["with_issues"] = sum(1 for d in details if d["issues"])

        with connect(self.db_path) as conn:
            conn.execute(
                "UPDATE upload_batches SET status = 'COMMITTED', rows = '[]', result = ?, committed_at = ? WHERE batch_id = ?",
                (json.dumps(result, ensure_ascii=False), now, batch_id),
            )
        return self.get_batch(batch_id)
